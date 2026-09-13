#!/usr/bin/env python3
"""KRX 표준서 인터페이스목록 → conf/krx_trcodes.toml 생성기.

    python tools/gen_krx_trcodes.py

trcode 사전만 만든다. **멀티캐스트 IP·포트는 생성하지 않는다** — 회선 배정이라
표준서에 적힌 값이 우리 회선에 그대로 오지 않는다. 배포 설정은 사람이 적는다.

Rust 가 아니라 파이썬인 이유: 표준서 버전이 올라갈 때만 한 번 도는 스크립트다.
`calamine` 을 워크스페이스 의존성으로 들이는 값을 못 한다. 결과 TOML 은 커밋되고,
그게 빌드가 보는 유일한 입력이다.

trcode 5바이트 = [데이터구분 2][정보구분+시장구분 3].
  B601F → B6(우선호가) + 01F(kospi200선물)
"""

from __future__ import annotations

import datetime as _dt
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl 이 필요하다: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
SPEC = ROOT / "documents/krx/접속표준서(정보분배-UDP_TCP실시간)_v1.341-배포용.xlsx"
SPEC_VERSION = "1.341"
OUT = ROOT / "conf/krx_trcodes.toml"

# 인터페이스목록 시트의 컬럼 (0-based). 헤더가 두 줄 병합이라 인덱스로 잡는다.
COL_ID, COL_NAME, COL_PERIOD, COL_TRCODE, COL_MARKET, COL_LENGTH = 1, 2, 6, 8, 10, 30
FIRST_DATA_ROW = 5

CODE_RE = re.compile(r"^[A-Z0-9]{5}$")
# "(파생A) DRV : G701F, G702F"  /  "          KNX : B001X"  /  "LK000"
LINE_RE = re.compile(r"^(?:\((?P<grp>[^)]+)\)\s*)?(?:(?P<mkt>[A-Z]{2,8})\s*:)?\s*(?P<codes>.*)$")


def toml_str(s: str) -> str:
    """TOML 기본 문자열. 개행은 공백으로 접는다 (엑셀 셀에 줄바꿈이 섞여 있다)."""
    s = " ".join(str(s).split())
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"') + '"'


def load_product_groups(wb) -> dict[str, str]:
    """별첨-정보구분코드: 3바이트 정보구분+시장구분 → 한글 이름."""
    ws = wb["별첨-정보구분코드"]
    out: dict[str, str] = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        code, name = row[5], row[4]
        if not code or not name:
            continue
        code = str(code).strip()
        if len(code) == 3:
            out.setdefault(code, " ".join(str(name).split()))
    return out


def parse_trcode_cell(cell: str) -> list[tuple[str, str | None, str | None]]:
    """'(파생A) DRV : G701F, G702F' → [(code, channel, market), ...].

    줄바꿈은 이어쓰기다: 괄호가 없으면 앞 줄의 회선을, 콜론이 없으면 앞 줄의 시장을
    물려받는다. 표준서가 셀 안에서 그렇게 접혀 있다.
    """
    out: list[tuple[str, str | None, str | None]] = []
    channel: str | None = None
    market: str | None = None
    for line in str(cell).split("\n"):
        line = line.strip()
        if not line or line == "-":
            continue
        m = LINE_RE.match(line)
        if not m:
            continue
        if m.group("grp"):
            channel = m.group("grp").strip()
        if m.group("mkt"):
            market = m.group("mkt").strip()
        for tok in re.split(r"[,/]", m.group("codes")):
            tok = tok.strip()
            if CODE_RE.match(tok):
                out.append((tok, channel, market))
    return out


def split_codes(
    seen: dict[str, list[dict]], interfaces: dict[str, dict], warnings: list[str]
) -> tuple[dict[str, dict], dict[str, dict]]:
    """trcode → 인터페이스를 확정 테이블과 회선별 테이블로 가른다.

    **trcode 는 인터페이스에 1:1 이 아니다.** 표준서에 세 가지 형태로 나온다:

    1. 회선마다 다른 전문 — `B201S` 는 증권A 회선에서 685B `증권 Snapshot`,
       주식파생 기초자산 회선에서 136B. 5바이트만 보고는 못 고른다.
       → `[code_by_channel]`. 회선을 모르면 디코드하면 안 된다.
    2. 같은 전문, 다른 시각 — `H101F` 는 IFMSSTD0005(장중 30초 주기)와
       IFMSRID0007(확정치 07:30/16:30) 양쪽에 있고 **길이가 같다**.
       레이아웃이 같으니 디코드는 하나로 되고, 의미만 시각으로 갈린다.
       → `[code]` + `also`.
    3. 전문 안의 구분자 — `M200G / 1`(시세) vs `/ 2`(종가). 길이 같음. 2와 같이 다룬다.
    """
    codes: dict[str, dict] = {}
    by_channel: dict[str, dict] = {}

    for code, entries in seen.items():
        uniq: list[dict] = []
        for e in entries:
            if e not in uniq:
                uniq.append(e)
        ids = []
        for e in uniq:
            if e["interface"] not in ids:
                ids.append(e["interface"])

        if len(ids) == 1:
            codes[code] = {
                "interface": ids[0],
                "group": code[2:],
                "channel": uniq[0]["channel"],
                "market": uniq[0]["market"],
                "also": [],
            }
            continue

        lengths = {interfaces[i]["length"] for i in ids}
        if len(lengths) == 1:
            # 같은 레이아웃. 디코더는 하나, 의미는 시각·전문 내용으로 갈린다.
            codes[code] = {
                "interface": ids[0],
                "group": code[2:],
                "channel": uniq[0]["channel"],
                "market": uniq[0]["market"],
                "also": ids[1:],
            }
            continue

        # 길이가 다르다 — 회선이 유일한 판별자다.
        mapping: dict[str, str] = {}
        for e in uniq:
            ch = e["channel"]
            if not ch:
                warnings.append(f"{code}: 길이가 갈리는데 회선이 비어 있다 ({e['interface']})")
                continue
            if ch in mapping and mapping[ch] != e["interface"]:
                warnings.append(
                    f"{code}: 같은 회선 {ch} 에 길이가 다른 전문 두 개 "
                    f"({mapping[ch]}, {e['interface']}) — 5바이트로 못 가른다"
                )
                continue
            mapping[ch] = e["interface"]
        by_channel[code] = {"group": code[2:], "by": mapping}

    return codes, by_channel


def main() -> int:
    if not SPEC.exists():
        sys.exit(f"표준서가 없다: {SPEC}")

    wb = openpyxl.load_workbook(SPEC, read_only=True, data_only=True)
    groups = load_product_groups(wb)
    ws = wb["인터페이스목록"]

    interfaces: dict[str, dict] = {}
    # trcode 는 인터페이스에 1:1 이 아니다 (아래 split_codes 주석 참조).
    seen: dict[str, list[dict]] = {}
    warnings: list[str] = []

    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        iid = row[COL_ID]
        if not iid or not str(iid).strip().startswith("IF"):
            continue
        iid = str(iid).strip()
        length = row[COL_LENGTH]
        if not isinstance(length, int):
            warnings.append(f"{iid}: 길이가 정수가 아니다 ({length!r}) — 건너뛴다")
            continue

        interfaces[iid] = {
            "name": " ".join(str(row[COL_NAME] or "").split()),
            "length": length,
            "period": " ".join(str(row[COL_PERIOD] or "").split()),
            "market": "/".join(str(row[COL_MARKET] or "").split()),
        }

        for code, channel, market in parse_trcode_cell(row[COL_TRCODE] or ""):
            if code[2:] not in groups:
                warnings.append(f"{code}: 상품군 {code[2:]} 이 별첨에 없다")
            seen.setdefault(code, []).append(
                {"interface": iid, "channel": channel or "", "market": market or ""}
            )

    codes, by_channel = split_codes(seen, interfaces, warnings)

    # 데이터구분 2바이트 → 그 코드를 쓰는 인터페이스들. 표준서에 표가 없어서 역산한다.
    data_class: dict[str, list[str]] = {}
    for code, meta in codes.items():
        for iid in [meta["interface"], *meta["also"]]:
            data_class.setdefault(code[:2], [])
            if iid not in data_class[code[:2]]:
                data_class[code[:2]].append(iid)
    for code, meta in by_channel.items():
        for iid in meta["by"].values():
            data_class.setdefault(code[:2], [])
            if iid not in data_class[code[:2]]:
                data_class[code[:2]].append(iid)

    today = _dt.date.today().isoformat()
    lines: list[str] = [
        "# GENERATED — 손으로 고치지 않는다. tools/gen_krx_trcodes.py 를 다시 돌린다.",
        f"# 원본: {SPEC.name}",
        "#",
        "# 여기에 **멀티캐스트 IP·포트는 없다.** 회선 배정이라 표준서 값이 우리 회선에",
        "# 그대로 오지 않는다. 배포 설정(conf/krx.toml)은 사람이 적고, 기동 시 거기 적힌",
        "# trcode 가 이 표에 있는지만 검증한다.",
        "",
        f'spec_version = {toml_str(SPEC_VERSION)}',
        f'generated = {toml_str(today)}',
        f"interface_count = {len(interfaces)}",
        f"code_count = {len(codes)}",
        f"ambiguous_code_count = {len(by_channel)}",
        "",
        "# 정보구분+시장구분 3바이트 (trcode 의 뒤 3자리) → 상품군.",
        "# 멀티캐스트 그룹이 상품군마다 갈리므로, 전략이 안 보는 상품군은 소켓을 안 연다.",
        "[product_group]",
    ]
    for code in sorted(groups):
        lines.append(f'"{code}" = {toml_str(groups[code])}')

    lines += [
        "",
        "# 데이터구분 2바이트 (trcode 의 앞 2자리) → 이 구분을 쓰는 인터페이스.",
        "# 표준서에 표가 없어서 코드에서 역산했다.",
        "[data_class]",
    ]
    for dc in sorted(data_class):
        ids = ", ".join(toml_str(i) for i in sorted(data_class[dc]))
        names = " / ".join(interfaces[i]["name"] for i in sorted(data_class[dc]))
        lines.append(f"# {dc}: {names}")
        lines.append(f'"{dc}" = [{ids}]')

    lines += [
        "",
        "# 인터페이스 → 전문 길이. 길이는 인터페이스 단위로 고정이다.",
        "# 디코더는 길이 + 종료키워드(0xFF)를 먼저 보고 그 다음에 필드를 읽는다.",
        "[interface]",
    ]
    for iid in sorted(interfaces):
        it = interfaces[iid]
        lines.append(
            f'{iid} = {{ length = {it["length"]}, market = {toml_str(it["market"])}, '
            f'period = {toml_str(it["period"])}, name = {toml_str(it["name"])} }}'
        )

    lines += [
        "",
        "# trcode → 인터페이스. 길이는 [interface] 에서 끌어온다 (중복 저장하지 않는다).",
        "# channel 은 표준서의 회선 그룹(증권A/파생A/…)이지 우리 회선이 아니다.",
        "[code]",
    ]
    for code in sorted(codes):
        c = codes[code]
        also = ""
        if c["also"]:
            ids = ", ".join(toml_str(i) for i in c["also"])
            also = f", also = [{ids}]"
        lines.append(
            f'{code} = {{ interface = {toml_str(c["interface"])}, '
            f'group = "{c["group"]}", channel = {toml_str(c["channel"])}, '
            f'market = {toml_str(c["market"])}{also} }}'
        )

    lines += [
        "",
        "# **같은 trcode 가 회선마다 다른 전문으로 온다.** 5바이트만 보고는 못 고른다 —",
        "# 예: B201S 는 증권A 회선에서 685B, 주식파생 기초자산 회선에서 136B.",
        "# 이 코드들은 [code] 에 없다. 어느 회선에서 받았는지를 같이 넘겨야 한다.",
        "[code_by_channel]",
    ]
    for code in sorted(by_channel):
        c = by_channel[code]
        pairs = ", ".join(f"{toml_str(k)} = {toml_str(v)}" for k, v in sorted(c["by"].items()))
        lines.append(f'{code} = {{ group = "{c["group"]}", by = {{ {pairs} }} }}')

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(
        f"{OUT.relative_to(ROOT)}: 인터페이스 {len(interfaces)}, "
        f"trcode {len(codes)} (+회선별 {len(by_channel)})"
    )
    for w in warnings:
        print(f"  경고: {w}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
